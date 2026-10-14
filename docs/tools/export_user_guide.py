"""Export the canonical illustrated Seis3D Markdown guide to an editable Word file."""

from __future__ import annotations

import argparse
import os
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
WORKSPACE_ROOT = ROOT.parent
SOURCE = ROOT / "docs/user/Seis3D_User_Guide.md"
DEFAULT_OUTPUT = WORKSPACE_ROOT / "out/docs/Seis3D_User_Guide.docx"
CONTENT_WIDTH = 6.7
INLINE_TOKEN = re.compile(r"(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")

def configure_page(section, illustrated: bool = False) -> None:
    """Use A4 prose and A3 landscape original screenshots with consistent margins."""
    section.orientation = WD_ORIENT.LANDSCAPE if illustrated else WD_ORIENT.PORTRAIT
    section.page_width = Cm(42 if illustrated else 21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5 if illustrated else 1.75)
    section.bottom_margin = Cm(1.5 if illustrated else 1.75)
    section.left_margin = Cm(1.5 if illustrated else 1.98)
    section.right_margin = Cm(1.5 if illustrated else 1.98)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.75)


def set_font(font, size: float, bold: bool = False, color: str = "000000") -> None:
    """Set Latin and Chinese faces explicitly for predictable Word rendering."""
    font.name = "Microsoft YaHei"
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = RGBColor.from_string(color)
    font._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")


def hyperlink(paragraph, label: str, target: str, output: Path) -> None:
    """Keep Markdown links clickable while preserving portable local paths."""
    if not re.match(r"^[a-zA-Z][\w+.-]*:", target):
        resolved = (SOURCE.parent / target).resolve()
        target = os.path.relpath(resolved, output.parent).replace("\\", "/")
    relationship = paragraph.part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "246B70")
    properties.append(color)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    element.append(run)
    paragraph._p.append(element)


def add_inline(paragraph, text: str, output: Path) -> None:
    """Translate the small inline Markdown subset used by the manual."""
    for token in INLINE_TOKEN.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif token.startswith("["):
            match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if match:
                hyperlink(paragraph, match[1], match[2], output)
            else:
                paragraph.add_run(token)
        else:
            paragraph.add_run(token)


def add_numbering(document, ordered: bool) -> int:
    """Create a fresh semantic list so numbered procedures restart at one."""
    numbering = document.part.numbering_part.element
    abstract_ids = [int(n.get(qn("w:abstractNumId"))) for n in numbering.findall(qn("w:abstractNum"))]
    number_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    number_id = max(number_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    single_level = OxmlElement("w:multiLevelType")
    single_level.set(qn("w:val"), "singleLevel")
    abstract.append(single_level)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("start", "1"), ("numFmt", "decimal" if ordered else "bullet"), ("lvlText", "%1." if ordered else "•"), ("lvlJc", "left")):
        item = OxmlElement(f"w:{tag}")
        item.set(qn("w:val"), value)
        level.append(item)
    properties = OxmlElement("w:pPr")
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "300")
    indent.set(qn("w:hanging"), "300")
    properties.append(indent)
    level.append(properties)
    abstract.append(level)
    numbering.insert_element_before(abstract, "w:num", "w:numIdMacAtCleanup")
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(number_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), str(abstract_id))
    number.append(reference)
    numbering.insert_element_before(number, "w:numIdMacAtCleanup")
    return number_id


def table_widths(headers: list[str]) -> list[float]:
    """Give short parameter fields less width than explanatory columns."""
    if len(headers) == 2:
        if headers[0] == "现象":
            return [2.25, 4.45]
        return [2.05, 4.65]
    if headers[0] == "区域":
        return [1.30, 2.90, 2.50]
    if headers[0] == "参数":
        return [2.15, 1.70, 2.85]
    if headers[0] == "物理网格":
        return [1.65, 2.30, 2.75]
    if headers[0] == "预览类型":
        return [1.65, 2.45, 2.60]
    return [CONTENT_WIDTH / len(headers)] * len(headers)


def add_table(document, lines: list[str], output: Path) -> None:
    """Build accessible, padded tables with repeat headers and flexible heights."""
    rows = [[value.strip() for value in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", cell) for cell in row)]
    widths = table_widths(rows[0])
    table = document.add_table(rows=0, cols=len(widths))
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            item.set(qn(f"w:{key}"), value)
        borders.append(item)
    properties.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, value in (("top", "85"), ("bottom", "85"), ("left", "105"), ("right", "105")):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:w"), value)
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    properties.append(margins)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        if row_index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for column_index, (cell, text, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E6EEF0" if row_index == 0 else "FFFFFF")
            cell._tc.get_or_add_tcPr().append(shading)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            paragraph.paragraph_format.keep_with_next = False
            if row_index and rows[0][0] in {"剖面", "物理网格", "快捷键"}:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(paragraph, text, output)
            for run in paragraph.runs:
                set_font(run.font, 9.3, bold=row_index == 0)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = Pt(3)
    spacer.paragraph_format.space_before = Pt(0)


def configure_document(document) -> None:
    """Apply the restrained Paper manual typography and page geometry."""
    section = document.sections[0]
    configure_page(section)
    normal = document.styles["Normal"]
    set_font(normal.font, 10.5)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in (("Title", 24), ("Subtitle", 10.5), ("Heading 1", 16), ("Heading 2", 12.5)):
        style = document.styles[name]
        set_font(style.font, size, bold=name not in {"Subtitle"})
        style.paragraph_format.space_before = Pt(10 if name.startswith("Heading") else 0)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True
        properties = style.element.get_or_add_pPr()
        for border in list(properties.findall(qn("w:pBdr"))):
            properties.remove(border)
    caption = document.styles["Caption"]
    set_font(caption.font, 9, color="525A59")
    caption.paragraph_format.line_spacing = 1.1
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False
    header = section.header.paragraphs[0]
    run = header.add_run("Seis3D 三维地震波模拟使用说明")
    set_font(run.font, 8)
    header.paragraph_format.space_after = Pt(0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Seis3D  |  ")
    set_font(run.font, 8, color="525A59")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.title = "Seis3D 三维地震波模拟使用说明"
    document.core_properties.subject = "浅色 Paper 界面下的运行 剖面交互与三维采集"
    document.core_properties.keywords = "Seis3D, CUDA, Paper, 三维弹性波, SEG-Y"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""


def export_guide(output: Path) -> tuple[int, int]:
    """Convert the maintained Markdown source without changing its technical content."""
    document = Document()
    configure_document(document)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    page_break = False
    list_number = None
    list_ordered = None
    images = 0
    chapters = 0
    illustrations: dict[Path, str] = {}
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line == "<!-- pagebreak -->":
            page_break = True
            list_number = None
            index += 1
            continue
        if line.startswith("# "):
            document.add_paragraph("SeisWave Lab · Seis3D", style="Title")
            document.add_paragraph("三维地震波模拟使用说明", style="Heading 1")
        elif line.startswith("## "):
            paragraph = document.add_paragraph(line[3:], style="Heading 1")
            paragraph.paragraph_format.page_break_before = page_break
            page_break = False
            chapters += 1
            list_number = None
        elif line.startswith("### "):
            document.add_paragraph(line[4:], style="Heading 2")
            list_number = None
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
            if not match:
                raise ValueError(f"Unsupported image syntax: {line}")
            path = (SOURCE.parent / match[2]).resolve()
            caption_text = match[1]
            next_index = index + 1
            while next_index < len(lines) and not lines[next_index].strip():
                next_index += 1
            if next_index < len(lines) and re.match(r"^图\d+\u3000", lines[next_index].strip()):
                caption_text = lines[next_index].strip()
                index = next_index
            caption = document.add_paragraph(style="Caption")
            add_inline(caption, f"截图{path.name[:2]}（完整窗口见附录）：{caption_text}", output)
            illustrations.setdefault(path, caption_text)
            images += 1
            list_number = None
        elif line.startswith("|"):
            block = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index].strip())
                index += 1
            add_table(document, block, output)
            list_number = None
            continue
        elif re.match(r"^(\d+\. |\- )", line):
            ordered = bool(re.match(r"^\d+\. ", line))
            if list_number is None or list_ordered != ordered:
                list_number = add_numbering(document, ordered)
                list_ordered = ordered
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(4)
            num_properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
            num_properties.get_or_add_ilvl().val = 0
            num_properties.get_or_add_numId().val = list_number
            add_inline(paragraph, re.sub(r"^(\d+\. |\- )", "", line), output)
        else:
            style = "Caption" if re.match(r"^图\d+\u3000", line) else "Normal"
            if line.startswith("适用项目 "):
                style = "Subtitle"
            paragraph = document.add_paragraph(style=style)
            if style == "Caption":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, line, output)
            list_number = None
        index += 1
    document.add_paragraph("附录：完整窗口截图", style="Heading 1")
    document.add_paragraph("截图按软件流程排列，每幅独占A3横向页。全部嵌入原始PNG，不裁切、不重采样、保持宽高比例；正文中的截图编号对应文件名前两位。")
    for path, caption_text in sorted(illustrations.items()):
        configure_page(document.add_section(WD_SECTION_START.NEW_PAGE), illustrated=True)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_with_next = True
        picture = paragraph.add_run().add_picture(str(path), width=Cm(39))
        picture._inline.docPr.set("descr", caption_text)
        caption = document.add_paragraph(style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(caption, f"截图{path.name[:2]}　{caption_text}", output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return chapters, len(illustrations)


def main() -> None:
    """Allow repeatable local export with all generated files outside the source tree."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    arguments = parser.parse_args()
    output = arguments.output.resolve()
    chapters, images = export_guide(output)
    print(f"Created {output}; {chapters} chapters and {images} embedded screenshots")


if __name__ == "__main__":
    main()
